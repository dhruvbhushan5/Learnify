from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Learnify_IP_Report_PreFinal.docx"
FONT_NAME = "Times New Roman"
BODY_SIZE = Pt(12)
ACCENT = RGBColor(31, 78, 121)
MUTED = RGBColor(95, 95, 95)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = FONT_NAME
    run.font.size = BODY_SIZE


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = FONT_NAME
    run.font.size = BODY_SIZE
    begin_run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(fld_begin)
    instr_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    instr_run._r.append(instr)
    end_run = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    end_run._r.append(fld_end)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = BODY_SIZE
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for name, size in [("Title", 22), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)]:
        style = styles[name]
        style.font.name = FONT_NAME
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT if name != "Title" else RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)


def add_header_footer(section):
    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("Learnify - Integrated Project Report")
    run.font.name = FONT_NAME
    run.font.size = BODY_SIZE
    run.font.color.rgb = MUTED
    p_pr = header._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D9D9D9")
    border.append(bottom)
    p_pr.append(border)
    add_page_number(section.footer.paragraphs[0])


def center_para(doc, text="", bold=False, size=12, color=None, after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = FONT_NAME
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    return p


def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0].cells
    for idx, title in enumerate(headers):
        set_cell_text(hdr[idx], title, True)
        set_cell_shading(hdr[idx], "EAF2F8")
        hdr[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], str(value))
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = width
    doc.add_paragraph()
    return table


def add_code(doc, code):
    for line in code.strip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line)
        r.font.name = FONT_NAME
        r.font.size = BODY_SIZE


def add_picture_if_exists(doc, path, caption, width=5.5):
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        cap = center_para(doc, caption, size=9, color=MUTED, after=10)
        cap.italic = True


def normalize_document_text(doc):
    paragraphs = list(doc.paragraphs)
    for section in doc.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)

    for paragraph in paragraphs:
        paragraph.paragraph_format.line_spacing = 1.5
        for run in paragraph.runs:
            run.font.name = FONT_NAME
            run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
            run.font.size = BODY_SIZE


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    configure_styles(doc)
    add_header_footer(section)

    center_para(doc, "Learnify - Online Learning Platform", True, 20)
    center_para(doc, "Integrated Project Report", True, 16, ACCENT, 18)
    center_para(doc, "Submitted in Partial Fulfilment of the requirement for the Course Integrated Project (22CS038) of", size=11)
    center_para(doc, "Computer Science and Engineering", True, 12)
    center_para(doc, "BE Batch - 2023", size=11)
    center_para(doc, "In", size=11)
    center_para(doc, "May 2026", True, 12, after=22)
    center_para(doc, "Under the Guidance of:                                      Submitted By:", True, 11)
    center_para(doc, "Dr. Gurtej Kaur                                      Aadeesh Jain, 2310991501 (G-17)", size=10)
    center_para(doc, "Associate Professor                                  Aakash Vij, 2310991502 (G-17)", size=10)
    center_para(doc, "                                                     Anmol, 2310991525 (G-17)", size=10)
    center_para(doc, "                                                     Garv Gupta, 2310991573 (G-17)", size=10, after=25)
    center_para(doc, "Department of Computer Science and Engineering", True, 11)
    center_para(doc, "Chitkara University Institute of Engineering & Technology,", size=11)
    center_para(doc, "Chitkara University, Punjab", size=11)

    doc.add_page_break()
    center_para(doc, "Certificate", True, 18, ACCENT, 18)
    add_body(doc, 'This is to certify that the project entitled "Learnify - Online Learning Platform" has been submitted for the Bachelor of Computer Science Engineering at Chitkara University, Punjab during the academic semester January 2026 to May 2026. The report has been prepared by Aadeesh Jain, Aakash Vij, Anmol, and Garv Gupta under my guidance and supervision.')
    add_body(doc, "The submitted work demonstrates the design and implementation of a web-based learning platform developed using Node.js, Express.js, EJS templates, MongoDB, Mongoose, and supporting middleware for security, routing, validation, and deployment.")
    doc.add_paragraph("\n\n")
    add_body(doc, "Dr. Gurtej Kaur                                                                 Associate Professor - CSE")

    doc.add_page_break()
    center_para(doc, "Candidate's Declaration", True, 18, ACCENT, 18)
    add_body(doc, 'We, Aadeesh Jain (2310991501), Aakash Vij (2310991502), Anmol (2310991525), and Garv Gupta (2310991573), B.E. - 2023 of Chitkara University, Punjab, hereby declare that the Integrated Project Report entitled "Learnify - Online Learning Platform" is our original work carried out as part of the Integrated Project course.')
    add_body(doc, "The report has not been submitted to any other university or institution for the award of any degree or diploma. All sources, tools, technologies, and references used during the development of the project have been acknowledged appropriately.")
    doc.add_paragraph("\n")
    center_para(doc, "Aadeesh Jain          Aakash Vij          Anmol          Garv Gupta", True, 10)
    center_para(doc, "2310991501            2310991502          2310991525     2310991573", size=10)

    doc.add_page_break()
    center_para(doc, "ACKNOWLEDGEMENT", True, 18, ACCENT, 16)
    add_body(doc, "It is our pleasure to express sincere gratitude to everyone who contributed directly or indirectly to the successful completion of this project. We are thankful to our guide, Dr. Gurtej Kaur, for continuous support, valuable suggestions, and guidance throughout the development of Learnify.")
    add_body(doc, "We also thank the Department of Computer Science and Engineering, Chitkara University, for providing the academic environment, technical exposure, and resources required to complete this project. Finally, we extend appreciation to our teammates for their coordination, research, implementation effort, and testing support.")

    doc.add_page_break()
    center_para(doc, "ABSTRACT", True, 18, ACCENT, 16)
    add_body(doc, "Learnify is a web-based online learning platform developed to provide students with an organized and accessible environment for browsing courses, viewing playlists, registering accounts, exploring teachers, and purchasing courses. The project focuses on creating a simple learning workflow where users can authenticate, browse educational content, and complete course purchase details through a structured interface.")
    add_body(doc, "The backend of the application is developed using Node.js and Express.js. Express routes manage page rendering, authentication, playlist operations, and course-purchase flows. EJS is used as the server-side view engine to render dynamic pages, while static assets such as CSS, images, thumbnails, and course media are served from the public directory.")
    add_body(doc, "MongoDB is used as the database with Mongoose models for users, videos, and purchases. User passwords are protected with bcrypt hashing through a Mongoose pre-save hook, improving security for stored credentials. The application also uses middleware such as Helmet, CORS, Morgan, compression, body-parser, and cookie-parser to support security, request logging, performance, and request parsing.")
    add_body(doc, "Learnify is designed for deployment on Render with environment-based configuration. The project demonstrates practical backend development, database connectivity, route organization, authentication, templating, and deployment practices in a full-stack web application.")

    doc.add_page_break()
    doc.add_heading("Table of Contents", level=1)
    add_numbered(doc, [
        "Abstract and Keywords",
        "Introduction to the Project: Background and Problem Statement",
        "Software and Hardware Requirement Specification: Methods, Programming Environment, and Requirements to Run the Application",
        "High Level and Low Level Design",
        "Database Analysis, Design and Implementation",
        "Program Structure and Database Connections",
        "Code Implementation and GUI Construction",
        "System Testing",
        "Conclusion",
        "Future Scope",
        "Bibliography and References",
    ])

    doc.add_page_break()
    doc.add_heading("Abstract and Keywords", level=1)
    add_body(doc, "Learnify is a full-stack online learning web application that provides core learning-platform features such as user registration, login, course listing, playlist viewing, teacher pages, course details, and payment-record storage. The system is implemented using Node.js, Express.js, EJS, MongoDB, and Mongoose, with static assets and templates organized into separate folders for maintainable development.")
    add_body(doc, "The project uses Express routing to separate page routes, authentication routes, playlist routes, and purchase routes. MongoDB stores persistent application data, while Mongoose schemas define the structure of users, videos, and purchases. The platform includes course data for HTML, CSS, JavaScript, Java, and C++ with tutors, prices, thumbnails, levels, descriptions, and video lists.")
    doc.add_heading("KEYWORDS", level=1)
    add_bullets(doc, ["Online Learning Platform", "Node.js", "Express.js", "EJS", "MongoDB", "Mongoose", "bcrypt", "Authentication", "Course Purchase", "Playlist Management", "Render Deployment", "Full Stack Development"])

    doc.add_page_break()
    doc.add_heading("2. Introduction to the Project", level=1)
    add_body(doc, "The rapid growth of digital education has increased the need for platforms that can organize learning material, present courses clearly, and support students through a simple web interface. Learnify is designed as an online learning system that combines course discovery, playlist-based video learning, authentication, teacher information, and purchase handling in one application.")
    add_body(doc, "The application follows a server-rendered full-stack architecture. Node.js provides the runtime environment, Express.js manages routing and middleware, EJS renders dynamic pages, and MongoDB stores user, playlist, and purchase data. The project is structured so that routes, models, views, public assets, and configuration files are separated clearly.")
    doc.add_heading("2.1 Background", level=2)
    add_body(doc, "Traditional learning resources are often scattered across multiple websites, files, and video platforms. Students may face difficulty finding structured course content, understanding the tutor profile, and accessing playlists in a single location. Learnify addresses this problem by presenting course content through a centralized learning interface.")
    add_body(doc, "The project uses modern JavaScript backend development practices. Server-side rendering with EJS keeps the interface lightweight, while MongoDB allows flexible storage for users, videos, and purchase records. Middleware improves logging, security headers, response compression, request parsing, and cross-origin handling.")
    doc.add_heading("2.2 Problem Statement", level=2)
    add_body(doc, "The problem addressed by Learnify is the lack of a simple and organized course platform for students to register, browse courses, access playlists, learn from tutor information, and record course purchases. The system must be easy to run locally, configurable for deployment, secure for basic authentication, and maintainable for future expansion.")

    doc.add_page_break()
    doc.add_heading("3. Software and Hardware Requirement Specification", level=1)
    doc.add_heading("3.1 Software Requirements", level=2)
    add_bullets(doc, [
        "Runtime: Node.js version 18 or above.",
        "Backend Framework: Express.js.",
        "Template Engine: EJS.",
        "Database: MongoDB local server or MongoDB Atlas.",
        "ODM: Mongoose.",
        "Authentication Security: bcryptjs password hashing.",
        "Middleware: helmet, cors, morgan, compression, body-parser, and cookie-parser.",
        "Deployment: Render web service with environment variables.",
        "Development Tools: Visual Studio Code, browser, terminal, Git, and GitHub.",
    ])
    doc.add_heading("3.2 Hardware Requirements", level=2)
    add_bullets(doc, [
        "Processor: Minimum dual-core processor; quad-core processor or better is recommended.",
        "RAM: Minimum 4 GB; 8 GB or more is recommended.",
        "Storage: Minimum 1 GB free space; 5 GB is recommended for project files, assets, and dependencies.",
        "Network: Internet is required for package installation and deployment; stable broadband is recommended.",
    ])
    doc.add_heading("3.3 Methods", level=2)
    add_bullets(doc, [
        "Requirement analysis was performed by identifying the student, course, playlist, teacher, and purchase flows.",
        "Backend development was completed using modular Express routers and Mongoose models.",
        "Database design was implemented through schemas for users, videos, and purchases.",
        "Interface construction was completed using EJS views and static CSS/images.",
        "Testing was carried out by running the server, checking routes, validating form input, and confirming database operations.",
    ])
    doc.add_heading("3.4 Programming Environment", level=2)
    add_body(doc, "The programming environment includes Node.js, npm, Express.js, EJS, MongoDB, and Mongoose. The application can be started with npm start after dependencies are installed and the MongoDB connection string is configured.")
    doc.add_heading("3.5 Requirements to Run the Application", level=2)
    add_code(doc, """
npm install
copy .env.example .env
npm run seed
npm start
""")
    add_body(doc, "The default local URL is http://localhost:8080. The .env file stores the PORT and MONGODB_URI values required by the server.")

    doc.add_page_break()
    doc.add_heading("4. High Level and Low Level Design", level=1)
    doc.add_heading("4.1 High Level Design", level=2)
    add_body(doc, "The high-level design of Learnify consists of four main layers: presentation layer, routing layer, business/data access layer, and database layer. The presentation layer contains EJS templates and public assets. The routing layer maps user requests to controllers implemented through Express routers. The data layer uses Mongoose models, and MongoDB stores persistent records.")
    add_bullets(doc, [
        "Presentation Layer: The views/*.ejs files and public/css and public/images folders display login, register, home, courses, playlist, teacher, and payment pages.",
        "Routing Layer: The routes/pages.js, routes/auth.js, routes/playlist.js, and routes/buyCourse.js files handle page navigation, authentication, playlist data, and course purchase requests.",
        "Model Layer: The models/User.js, models/Video.js, and models/Purchase.js files define database schemas and helper methods.",
        "Database Layer: The db.js file connects the application to MongoDB, where user, video, and purchase records are stored.",
    ])
    doc.add_heading("4.2 Low Level Design", level=2)
    add_body(doc, "At the low level, each feature is separated into route handlers and supporting models. The authentication module validates email and password fields, checks existing users, hashes passwords through the User schema, and compares passwords during login. The playlist module fetches videos from MongoDB and renders playlist pages. The purchase module validates payment form fields and stores purchase records.")
    add_bullets(doc, [
        "Registration Module: Takes email and password, validates the fields, checks duplicate email, hashes the password, saves the user, and shows a success message.",
        "Login Module: Takes email and password, finds the user, compares the password hash, and opens the home page or shows an error message.",
        "Playlist Module: Handles playlist requests, fetches or inserts videos using the Video model, and renders the playlist page.",
        "Course Purchase Module: Takes student name, email, phone, and payment method, validates the form, creates a Purchase document, and shows the payment success page.",
    ])

    doc.add_page_break()
    doc.add_heading("5. Database Analysis, Design and Implementation", level=1)
    add_body(doc, "MongoDB is used as the database for Learnify. It is suitable for this project because the data model is flexible and can store documents for users, videos, purchases, and course-related records. Mongoose is used to define schemas, validation rules, and model methods.")
    doc.add_heading("5.1 Main Collections", level=2)
    add_bullets(doc, [
        "users: Stores registered user accounts with email and hashed password fields.",
        "videos: Stores playlist content using title, videoUrl, and thumbnail fields.",
        "purchases: Stores course purchase records using courseSlug, courseTitle, amount, studentName, email, phone, paymentMethod, status, and timestamps.",
    ])
    doc.add_heading("5.2 Relationship Design", level=2)
    add_body(doc, "The purchase collection stores the selected course through courseSlug and courseTitle. This allows the purchase record to remain readable even if the static course list changes later. The video collection is independent and can be expanded for multiple playlists or course-specific mapping in future versions.")
    doc.add_heading("5.3 Security and Data Integrity", level=2)
    add_bullets(doc, [
        "User email is required and unique in the User schema.",
        "Passwords are hashed before saving through a Mongoose pre-save hook.",
        "Purchase payment method is restricted to upi, card, or netbanking.",
        "Purchase status is restricted to paid, pending, or failed.",
        "Required fields prevent incomplete user, video, and purchase documents.",
    ])

    doc.add_page_break()
    doc.add_heading("6. Program Structure and Database Connections", level=1)
    add_body(doc, "The project structure separates backend entry point, database configuration, routes, models, templates, public assets, and seed data. This organization makes the codebase easier to understand, maintain, and deploy.")
    add_bullets(doc, [
        "server.js: Main entry point that configures Express, middleware, routes, static files, and starts the server.",
        "db.js: Connects the application to MongoDB using Mongoose and MONGODB_URI.",
        "routes/: Contains page, authentication, playlist, and purchase route files.",
        "models/: Contains User, Video, and Purchase Mongoose schemas.",
        "views/: Contains EJS templates rendered by the server.",
        "public/: Contains CSS, HTML assets, images, thumbnails, and videos.",
        "data/courses.js: Stores static course metadata and a helper function for lookup by slug.",
        "seedVideos.js: Seeds initial video records into MongoDB.",
    ])
    doc.add_heading("6.1 Database Connection", level=2)
    add_code(doc, """
require("dotenv").config();
require("./db");
app.listen(PORT, "0.0.0.0", () => {
    console.log(`Server running at http://localhost:${PORT}`);
});
""")
    add_body(doc, "The application reads configuration from environment variables. For local use, .env can contain PORT=8080 and MONGODB_URI=mongodb://127.0.0.1:27017/learnifyDB. For deployment, Render supplies PORT automatically and MongoDB Atlas can supply the production database connection string.")

    doc.add_page_break()
    doc.add_heading("7. Code Implementation and GUI Construction", level=1)
    doc.add_heading("7.1 Backend Code Implementation", level=2)
    add_body(doc, "The backend implementation is based on Express routers. server.js registers common middleware, configures EJS, serves static files, connects routes, and attaches the error handler. The routes folder implements page navigation, login/register operations, playlist operations, and course-purchase logic.")
    add_code(doc, """
app.use("/", pageRoutes);
app.use("/auth", authRoutes);
app.use("/playlist", playlistRoutes);
app.use("/buy-course", buyCourseRoutes);
app.use(express.static(path.join(__dirname, "public")));
""")
    doc.add_heading("7.2 Authentication Implementation", level=2)
    add_body(doc, "The User schema stores email and password. Before saving a new user, the password is hashed with bcryptjs. During login, the entered password is compared with the stored hash through a schema method. This improves security by avoiding plain-text password storage.")
    add_code(doc, """
userSchema.pre("save", async function(next) {
  if (this.isModified("password")) {
    this.password = await bcrypt.hash(this.password, 10);
  }
  next();
});
""")
    doc.add_heading("7.3 GUI Construction", level=2)
    add_body(doc, "The graphical interface is built using EJS templates and static CSS. Important views include login, register, index/home, about, courses, playlist, buy-course, course-payment, payment-success, teachers, teacher-register, and contact pages. Course thumbnails and profile images are stored inside public/images.")
    add_picture_if_exists(doc, ROOT / "public" / "images" / "thumb-1.png", "Figure: Course thumbnail asset used in Learnify.", 3.5)
    add_picture_if_exists(doc, ROOT / "public" / "images" / "html-thumbnail.png", "Figure: HTML course visual asset used by the platform.", 3.5)
    doc.add_heading("7.4 Course Purchase Flow", level=2)
    add_body(doc, "The buy-course route displays available courses, opens a course-specific payment page using the course slug, validates student and payment details, creates a Purchase document, and renders the payment-success view. This demonstrates complete form handling and database persistence.")

    doc.add_page_break()
    doc.add_heading("8. System Testing", level=1)
    add_body(doc, "System testing verifies that all modules of Learnify work according to requirements. Testing focuses on route behavior, form validation, database operations, user authentication, static asset loading, EJS rendering, and deployment readiness.")
    doc.add_heading("8.1 Types of Testing Performed", level=2)
    add_bullets(doc, [
        "Unit Testing: Individual route handlers and model validations are checked to confirm that each module performs its function correctly.",
        "Integration Testing: Routes are tested with database operations and EJS rendering to confirm that data moves correctly between forms, routes, models, and views.",
        "Functional Testing: Login, register, courses, playlists, and payment flows are tested to confirm that users receive the correct page or validation message.",
        "Database Testing: MongoDB records for users, videos, and purchases are verified to confirm that documents are inserted and retrieved correctly.",
        "Security Testing: Password hashing and required-field validation are checked to confirm that passwords are stored as hashes and incomplete forms are rejected.",
        "Compatibility Testing: Application pages are checked in modern browsers to confirm that pages render correctly and static assets load.",
    ])
    doc.add_heading("8.2 Sample Test Cases", level=2)
    add_bullets(doc, [
        "Register New User: When a unique email and password are entered, the user is saved and the login page shows a success message.",
        "Register Duplicate User: When an existing email is entered, the register page shows a duplicate-user error.",
        "Login Valid User: When the correct email and password are entered, the home page is rendered.",
        "Login Invalid User: When the wrong email or password is entered, the login page shows an invalid credential error.",
        "Open Playlist: When /playlist is opened, videos are fetched from MongoDB and the playlist page is rendered.",
        "Buy Course: When the payment form is completed, the purchase is saved and the success page is rendered.",
    ])

    doc.add_page_break()
    doc.add_heading("9. Conclusion", level=1)
    add_body(doc, "Learnify successfully demonstrates the design and implementation of a full-stack online learning platform using Node.js, Express.js, EJS, MongoDB, and Mongoose. The application provides structured course browsing, authentication, playlist management, teacher-related pages, and course-purchase recording.")
    add_body(doc, "The project helped in understanding server-side rendering, Express routing, middleware usage, MongoDB integration, Mongoose schema design, password hashing, static asset management, and cloud deployment preparation. Learnify can be extended further into a production-grade learning management system.")

    doc.add_heading("10. Future Scope", level=1)
    add_bullets(doc, [
        "Add JWT or session-based protected routes for authenticated-only course access.",
        "Create separate student, teacher, and admin dashboards.",
        "Map videos directly to courses and track lesson completion progress.",
        "Integrate a real payment gateway and invoice generation.",
        "Add course reviews, ratings, certificates, and quiz modules.",
        "Improve responsive design and accessibility across all pages.",
        "Add automated tests for routes, models, and validation logic.",
        "Deploy with MongoDB Atlas, monitoring, logging, and environment-specific configuration.",
    ])

    doc.add_page_break()
    doc.add_heading("11. Bibliography and References", level=1)
    add_numbered(doc, [
        "Node.js Official Documentation - https://nodejs.org/en/docs",
        "Express.js Official Documentation - https://expressjs.com/",
        "MongoDB Documentation - https://www.mongodb.com/docs/",
        "Mongoose Documentation - https://mongoosejs.com/docs/",
        "EJS Documentation - https://ejs.co/",
        "bcrypt.js Package Documentation - https://www.npmjs.com/package/bcryptjs",
        "Render Deployment Documentation - https://render.com/docs",
        "Jon Duckett, HTML and CSS: Design and Build Websites, Wiley Publications.",
        "Elmasri and Navathe, Fundamentals of Database Systems, Pearson Education.",
    ])

    doc.core_properties.title = "Learnify - Integrated Project Report"
    doc.core_properties.subject = "Online Learning Platform"
    doc.core_properties.author = "Aadeesh Jain, Aakash Vij, Anmol, Garv Gupta"
    normalize_document_text(doc)
    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT)
